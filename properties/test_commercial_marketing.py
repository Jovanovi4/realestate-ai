from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth.models import User
from django.db import connection
from django.db.migrations.executor import MigrationExecutor
from django.test import TestCase, TransactionTestCase
from django.template.loader import render_to_string
from django.urls import reverse
from docx import Document
from reportlab.platypus import Paragraph, Table

from .models import Lead, Property
from .services.ai_service import AIService
from .services.presentation_docx_service import PresentationDocxService
from .services.presentation_service import PresentationService
from .services.property_marketing import format_money, price_summary, property_fact_rows
from .services.telegram_public_service import listing_text


class CommercialMarketingTests(TestCase):
    def setUp(self):
        owner = User.objects.create_user("commercial-marketing-owner", password="password")
        self.property = Property.objects.create(
            owner=owner, title="Помещение для салона", property_type="retail", deal_type="rent",
            price=Decimal("180000.00"), currency="RUB", area=Decimal("60.00"),
            commercial_purpose="Салон красоты", ceiling_height_m=Decimal("3.25"),
            power_kw=Decimal("25.50"), entrance_type="separate", parking_type="on_site",
            access_mode="round_the_clock", rooms=4, land_area=100, furnished=True,
        )

    def test_shared_price_and_facts_distinguish_whole_property_from_sqm_rate(self):
        rent = price_summary(self.property)
        self.assertEqual(rent.main, "180 000 ₽ / мес.")
        self.assertEqual(rent.scope, "Аренда всего помещения за месяц")
        self.assertEqual(rent.per_sqm, "Расчётно 3 000 ₽ / м² / мес.")
        facts = dict(property_fact_rows(self.property))
        self.assertEqual(facts["Площадь"], "60 м²")
        self.assertEqual(facts["Назначение"], "Салон красоты")
        self.assertEqual(facts["Высота потолков"], "3,25 м")
        self.assertEqual(facts["Мощность"], "25,5 кВт")
        self.assertEqual(facts["Вход"], "Отдельный вход")
        self.assertEqual(facts["Парковка"], "На территории")
        self.assertEqual(facts["Режим доступа"], "Круглосуточно")
        for residential in ("Комнаты", "Санузлы", "Участок", "Мебель"):
            self.assertNotIn(residential, facts)

        self.property.deal_type = "sale"
        self.property.price = Decimal("18000000.00")
        self.assertEqual(price_summary(self.property).scope, "Полная цена помещения")
        self.assertEqual(price_summary(self.property).per_sqm, "Расчётно 300 000 ₽ / м²")

    def test_price_has_one_format_in_crm_and_landing(self):
        self.assertEqual(format_money(Decimal("140000.00"), "RUB"), "140 000 ₽")
        self.assertEqual(format_money(Decimal("125000.50"), "USD"), "125 000,5 $")
        self.client.force_login(self.property.owner)
        for name in ("property_list", "property_detail"):
            args = [self.property.pk] if name == "property_detail" else []
            response = self.client.get(reverse(name, args=args))
            self.assertContains(response, "180 000 ₽ / мес.")
            self.assertNotContains(response, "180000,00")
        landing_price = render_to_string("properties/includes/landing_price.html", {"property": self.property})
        self.assertIn("180 000 ₽ / мес.", landing_price)
        self.assertIn("Расчётно 3 000 ₽ / м² / мес.", landing_price)

    def test_telegram_announcement_uses_commercial_units_and_facts(self):
        text = listing_text(self.property)
        for expected in (
            "180 000 ₽ / мес.", "Аренда всего помещения за месяц", "3 000 ₽ / м² / мес.",
            "Назначение: Салон красоты", "Мощность: 25,5 кВт", "Режим доступа: Круглосуточно",
        ):
            self.assertIn(expected, text)
        self.assertNotIn("комн.", text)
        self.assertLessEqual(len(text), 950)

    def test_pdf_and_docx_presentations_use_the_same_commercial_details(self):
        captured = []
        with patch("properties.services.presentation_service.SimpleDocTemplate.build", side_effect=lambda story, **kwargs: captured.extend(story)):
            PresentationService.build_pdf(self.property, None, include_benefits=False, include_contacts=False)
        pdf_text = []
        for flowable in captured:
            if isinstance(flowable, Paragraph):
                pdf_text.append(flowable.getPlainText())
            elif isinstance(flowable, Table):
                pdf_text.extend(cell.getPlainText() for row in flowable._cellvalues for cell in row if isinstance(cell, Paragraph))
        pdf_text = " ".join(pdf_text)

        docx = Document(BytesIO(PresentationDocxService.build_docx(
            self.property, None, include_benefits=False, include_contacts=False,
        )))
        docx_text = " ".join(
            [paragraph.text for paragraph in docx.paragraphs]
            + [cell.text for table in docx.tables for row in table.rows for cell in row.cells]
        )
        for content in (pdf_text, docx_text):
            for expected in (
                "180 000 ₽ / мес.", "Аренда всего помещения за месяц", "Расчётно 3 000 ₽ / м² / мес.",
                "Салон красоты", "3,25 м", "25,5 кВт", "Отдельный вход", "На территории", "Круглосуточно",
            ):
                self.assertIn(expected, content)
            for residential in ("Комнаты", "Участок", "Мебель"):
                self.assertNotIn(residential, content)

    @patch("properties.services.ai_service.AIService._run_prompt")
    def test_commercial_prompts_keep_living_language_out_of_all_generation_paths(self, run_prompt):
        prompt = AIService.build_prompt(self.property, "full_description", "business")
        self.assertIn("аренда коммерческого помещения", prompt)
        self.assertIn("Назначение: Салон красоты", prompt)
        self.assertIn("Режим доступа: Круглосуточно", prompt)
        self.assertNotIn("условия проживания", prompt)
        self.assertNotIn("- Комнаты:", prompt)
        self.assertNotIn("- Мебель:", prompt)

        run_prompt.return_value = {"content": "Уточните условия помещения."}
        AIService.generate_inline_content(self.property, "landing_contact_title", "business")
        self.assertIn("коммерческого помещения", run_prompt.call_args.args[0])
        self.assertNotIn("условия проживания", run_prompt.call_args.args[0])

        types = (
            "headline", "short_description", "full_description", "landing_headline", "landing_subtitle",
            "seo_title", "seo_description", "benefits", "cta",
        )
        run_prompt.return_value = {"content": "\n".join(f"[[{kind}]] Текст {kind}" for kind in types)}
        self.assertEqual(len(AIService.generate_bundle(self.property, "business", "package")), len(types))
        self.assertIn("коммерческого помещения", run_prompt.call_args.args[0])
        self.assertNotIn("условия проживания", run_prompt.call_args.args[0])

        run_prompt.return_value = {"content": "Здравствуйте! Уточним условия."}
        AIService.generate_lead_reply(self.property, Lead(name="Клиент", message="Интересует помещение"), "business")
        self.assertIn("коммерческого помещения", run_prompt.call_args.args[0])


class CommercialPurposeSearchMigrationTests(TransactionTestCase):
    def test_existing_commercial_purpose_is_normalized(self):
        previous = [("properties", "0050_client_commercial_search")]
        current = [("properties", "0051_commercial_purpose_search")]
        executor = MigrationExecutor(connection)
        latest = executor.loader.graph.leaf_nodes()
        self.addCleanup(lambda: MigrationExecutor(connection).migrate(latest))
        executor.migrate(previous)
        old_apps = executor.loader.project_state(previous).apps
        OldUser = old_apps.get_model("auth", "User")
        OldProperty = old_apps.get_model("properties", "Property")
        owner = OldUser.objects.create(username="commercial-search-migration-owner")
        property = OldProperty.objects.create(
            owner_id=owner.pk, title="Помещение", property_type="retail",
            commercial_purpose="Салон КРАСОТЫ",
        )

        MigrationExecutor(connection).migrate(current)
        new_apps = MigrationExecutor(connection).loader.project_state(current).apps
        NewProperty = new_apps.get_model("properties", "Property")
        self.assertEqual(
            NewProperty.objects.get(pk=property.pk).commercial_purpose_search,
            "салон красоты",
        )
