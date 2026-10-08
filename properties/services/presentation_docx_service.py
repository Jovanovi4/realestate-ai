from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor
from PIL import Image as PilImage

from .property_marketing import price_summary, property_fact_rows


class PresentationDocxService:
    @staticmethod
    def build_docx(property, profile, *, include_facts=True, include_description=True, include_benefits=True, include_contacts=True, photo_limit=6):
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Cm(1.7)
        section.left_margin = section.right_margin = Cm(1.8)
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10.5)

        title = document.add_heading(property.marketing_headline or property.title, 0)
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        meta = document.add_paragraph(f"{property.get_deal_type_display()} · {property.get_property_type_display()}")
        meta.style = document.styles["Subtitle"]
        price_details = price_summary(property)
        if price_details:
            price = document.add_paragraph()
            price_run = price.add_run(price_details.main)
            price_run.font.size = Pt(16)
            price_run.font.color.rgb = RGBColor(7, 134, 110)
            for line in (price_details.scope, price_details.per_sqm):
                if line:
                    document.add_paragraph(line, style="Subtitle")
        if property.address:
            document.add_paragraph(property.address)

        def add_image(property_image, width):
            try:
                property_image.image.open("rb")
                content = property_image.image.read()
                property_image.image.close()
                with PilImage.open(BytesIO(content)) as source:
                    source.thumbnail((1600, 1000), PilImage.Resampling.LANCZOS)
                    if source.mode in {"RGBA", "LA"}:
                        background = PilImage.new("RGB", source.size, "white")
                        background.paste(source, mask=source.getchannel("A"))
                        source = background
                    elif source.mode != "RGB":
                        source = source.convert("RGB")
                    image = BytesIO()
                    source.save(image, format="JPEG", quality=82, optimize=True)
                image.seek(0)
                document.add_picture(image, width=width)
            except OSError:
                return

        if property.primary_image:
            add_image(property.primary_image, Cm(16.5))
        if include_facts:
            facts = property_fact_rows(property)
            if facts:
                document.add_heading("Характеристики", level=1)
                table = document.add_table(rows=0, cols=2)
                table.style = "Light Shading Accent 1"
                for label, value in facts:
                    cells = table.add_row().cells
                    cells[0].text, cells[1].text = label, str(value)
        if include_description and (property.description or property.short_description):
            document.add_heading("Об объекте", level=1)
            document.add_paragraph(property.description or property.short_description)
        if include_benefits and property.landing_benefits:
            document.add_heading("Преимущества", level=1)
            for benefit in property.landing_benefits:
                if benefit.get("title"):
                    document.add_paragraph(f"{benefit['title']}. {benefit.get('description', '')}", style="List Bullet")
        images = list(property.images.all()[:photo_limit])
        if len(images) > 1:
            document.add_heading("Фотографии", level=1)
            for image in images[1:]:
                add_image(image, Cm(16.5))
        if include_contacts and profile:
            document.add_heading("Контакты риелтора", level=1)
            name = profile.display_name or property.owner.get_full_name() or property.owner.username
            contacts = ", ".join(filter(None, [profile.phone, f"Telegram: @{profile.telegram_username}" if profile.telegram_username else "", profile.email]))
            document.add_paragraph(name)
            if contacts:
                document.add_paragraph(contacts)
        output = BytesIO()
        document.save(output)
        return output.getvalue()
